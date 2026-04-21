from datetime import datetime
from pathlib import Path
from openpyxl import Workbook, load_workbook

FICHEIRO_EXCEL = "simulacao_medicoes.xlsx"


def criar_excel_se_nao_existir():
    caminho = Path(FICHEIRO_EXCEL)

    if caminho.exists():
        return

    wb = Workbook()
    ws = wb.active
    ws.title = "medicao_factos"

    ws.append([
        "id_medicao",
        "id_pedido",
        "source",
        "station_name",
        "station_id",
        "station_distanceKm",
        "station_latitude",
        "station_longitude",
        "station_elevationM",
        "requested_latitude",
        "requested_longitude",
        "observation_dataHora",
        "variavel",
        "valor",
        "timestamp_pedido"
    ])

    wb.save(FICHEIRO_EXCEL)


def proximo_id(ws, coluna_indice):
    if ws.max_row <= 1:
        return 1

    valores = []

    for row in ws.iter_rows(min_row=2, values_only=True):
        valor = row[coluna_indice]
        if valor is not None:
            valores.append(valor)

    if not valores:
        return 1

    return max(valores) + 1


def guardar_json_normalizado_em_excel(dados: dict):
    criar_excel_se_nao_existir()

    wb = load_workbook(FICHEIRO_EXCEL)
    ws = wb["medicao_factos"]

    id_pedido = proximo_id(ws, 1)  # coluna 2 = id_pedido
    timestamp_pedido = datetime.now().isoformat(timespec="seconds")

    requested_location = dados.get("requested_location", {})
    req_lat = requested_location.get("latitude")
    req_lon = requested_location.get("longitude")

    for source_key in ["foreca", "ipma"]:
        bloco_fonte = dados.get(source_key)

        if not bloco_fonte:
            continue

        source = bloco_fonte.get("source")
        station = bloco_fonte.get("station", {})
        observation = bloco_fonte.get("observation", {})

        station_name = station.get("name")
        station_id = station.get("id")
        station_distance_km = station.get("distanceKm")
        station_lat = station.get("latitude")
        station_lon = station.get("longitude")
        station_elev = station.get("elevationM")
        observation_data_hora = observation.get("dataHora")

        for variavel, valor in observation.items():
            id_medicao = proximo_id(ws, 0)  # coluna 1 = id_medicao

            ws.append([
                id_medicao,
                id_pedido,
                source,
                station_name,
                station_id,
                station_distance_km,
                station_lat,
                station_lon,
                station_elev,
                req_lat,
                req_lon,
                observation_data_hora,
                variavel,
                valor,
                timestamp_pedido
            ])

    wb.save(FICHEIRO_EXCEL)